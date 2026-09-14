from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "artifacts" / "tmp" / "TCC_Leonardo_Caseri_Rascunho_3_2_base.docx"
INTRO = ROOT / "tcc_texto" / "INTRODUCAO_PARCIAL.md"
THEORY = ROOT / "tcc_texto" / "FUNDAMENTACAO_TEORICA_PARCIAL.md"
DEVELOPMENT = ROOT / "tcc_texto" / "DESENVOLVIMENTO_PARCIAL.md"
SUMMARY = ROOT / "tcc_texto" / "RESUMO_PARCIAL.md"
ARCHITECTURE_FIGURE = ROOT / "artifacts" / "figures" / "architecture.png"
DASHBOARD_FIGURE = ROOT / "artifacts" / "figures" / "dashboard-overview.png"


def sections(path: Path) -> dict[str, str]:
    text = path.read_text(encoding="utf-8")
    result: dict[str, str] = {}
    current: str | None = None
    lines: list[str] = []
    for line in text.splitlines():
        if line.startswith("## "):
            if current is not None:
                result[current] = "\n".join(lines).strip()
            current = line[3:].strip()
            lines = []
        elif current is not None:
            lines.append(line)
    if current is not None:
        result[current] = "\n".join(lines).strip()
    return result


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D9D9D9", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, value: int = 80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge in ("top", "left", "bottom", "right"):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_no_wrap(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    if tc_pr.find(qn("w:noWrap")) is None:
        tc_pr.append(OxmlElement("w:noWrap"))


def mark_row_as_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)


def set_picture_alt_text(paragraph, description: str) -> None:
    for doc_pr in paragraph._p.xpath(".//wp:docPr"):
        doc_pr.set("descr", description)


def add_figure(doc: Document, path: Path, caption: str, description: str) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run().add_picture(str(path), width=Cm(15.3))
    set_picture_alt_text(paragraph, description)

    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.first_line_indent = Cm(0)
    caption_paragraph.paragraph_format.line_spacing = 1.0
    caption_paragraph.paragraph_format.space_after = Pt(0)
    caption_run = caption_paragraph.add_run(caption)
    caption_run.font.name = "Times New Roman"
    caption_run.font.size = Pt(10)

    source = doc.add_paragraph()
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source.paragraph_format.first_line_indent = Cm(0)
    source.paragraph_format.line_spacing = 1.0
    source.paragraph_format.space_after = Pt(6)
    source_run = source.add_run("Fonte: elaboração própria (2026).")
    source_run.font.name = "Times New Roman"
    source_run.font.size = Pt(10)


def clean_inline(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    return text.replace("—", "-").strip()


def add_markdown_block(doc: Document, block: str) -> None:
    lines = block.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line or line.startswith(">"):
            index += 1
            continue
        if line.startswith("| ") and index + 1 < len(lines) and "---" in lines[index + 1]:
            table_lines = [line]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = [[clean_inline(value.strip()) for value in row.strip("|").split("|")] for row in table_lines]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            mark_row_as_header(table.rows[0])
            width_sets = {
                2: [Cm(9), Cm(5.2)],
                3: [Cm(6), Cm(4), Cm(4.2)],
                5: [Cm(5), Cm(2.2), Cm(2.2), Cm(2.4), Cm(2.4)],
                6: [Cm(4.3), Cm(1.8), Cm(1.8), Cm(1.8), Cm(2.2), Cm(2.2)],
            }
            preferred_widths = width_sets.get(len(rows[0]), [Cm(14.2 / len(rows[0]))] * len(rows[0]))
            for col_index, width in enumerate(preferred_widths[: len(rows[0])]):
                table.columns[col_index].width = width
            for row_index, values in enumerate(rows):
                for col_index, value in enumerate(values):
                    cell = table.cell(row_index, col_index)
                    if col_index < len(preferred_widths):
                        cell.width = preferred_widths[col_index]
                    cell.text = value
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    set_cell_border(cell)
                    set_cell_margins(cell)
                    if col_index > 0:
                        set_cell_no_wrap(cell)
                    if row_index == 0:
                        set_cell_shading(cell, "1F4E5F")
                    elif row_index % 2 == 0:
                        set_cell_shading(cell, "F2F6F7")
                    for paragraph in cell.paragraphs:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for run in paragraph.runs:
                            run.font.name = "Times New Roman"
                            run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
                            run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
                            run.font.size = Pt(8.5)
                            if row_index == 0:
                                run.bold = True
                                run.font.color.rgb = RGBColor(255, 255, 255)
            doc.add_paragraph()
            continue
        if line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(clean_inline(line[2:]))
            index += 1
            continue
        paragraph_lines = [line]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or candidate.startswith("- ") or candidate.startswith("| "):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = doc.add_paragraph(clean_inline(" ".join(paragraph_lines)))
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)

    for style_name, size in (("Title", 14), ("Heading 1", 12), ("Heading 2", 12)):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["Heading 1"].paragraph_format.page_break_before = True
    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    if "Reference" not in [style.name for style in doc.styles]:
        reference = doc.styles.add_style("Reference", WD_STYLE_TYPE.PARAGRAPH)
    else:
        reference = doc.styles["Reference"]
    reference.font.name = "Times New Roman"
    reference._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    reference._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    reference.font.size = Pt(12)
    reference.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    reference.paragraph_format.first_line_indent = Cm(0)
    reference.paragraph_format.line_spacing = 1.0
    reference.paragraph_format.space_after = Pt(12)

    for footer in {section.footer for section in doc.sections}:
        add_page_number(footer.paragraphs[0])

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(54)
    run = p.add_run("UNIVERSIDADE DE ARARAQUARA\nUNIARA\nCURSO DE SISTEMAS DE INFORMAÇÃO")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(36)
    title.paragraph_format.space_after = Pt(54)
    direct_border = title._p.get_or_add_pPr().find(qn("w:pBdr"))
    if direct_border is not None:
        title._p.get_or_add_pPr().remove(direct_border)
    title.add_run("SISTEMA INTELIGENTE PARA ANÁLISE DE VENDAS E PRIORIZAÇÃO DE OPORTUNIDADES B2B")

    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.first_line_indent = Cm(0)
    author.add_run("Leonardo Caseri").bold = True

    advisor = doc.add_paragraph()
    advisor.alignment = WD_ALIGN_PARAGRAPH.CENTER
    advisor.paragraph_format.first_line_indent = Cm(0)
    advisor.add_run("Orientador: Felipe Diniz Dalilo")

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(120)

    city = doc.add_paragraph()
    city.alignment = WD_ALIGN_PARAGRAPH.CENTER
    city.paragraph_format.first_line_indent = Cm(0)
    city.add_run("Araraquara\n2026")
    doc.add_page_break()


def add_front_matter(doc: Document, summary: dict[str, str]) -> None:
    for title, keywords_label in (("Resumo", "Palavras-chave"), ("Abstract", "Keywords")):
        heading = doc.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.paragraph_format.first_line_indent = Cm(0)
        heading.paragraph_format.space_after = Pt(18)
        run = heading.add_run(title.upper())
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)

        body, keywords = summary[title].rsplit(f"\n\n{keywords_label}:", 1)
        paragraph = doc.add_paragraph(clean_inline(body))
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.first_line_indent = Cm(0)

        keywords_paragraph = doc.add_paragraph()
        keywords_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        keywords_paragraph.paragraph_format.first_line_indent = Cm(0)
        keywords_paragraph.add_run(f"{keywords_label}: ").bold = True
        keywords_paragraph.add_run(keywords.strip())
        doc.add_page_break()

    toc_heading = doc.add_paragraph()
    toc_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_heading.paragraph_format.first_line_indent = Cm(0)
    toc_heading.paragraph_format.space_after = Pt(18)
    toc_run = toc_heading.add_run("SUMÁRIO")
    toc_run.bold = True
    toc_run.font.name = "Times New Roman"
    toc_run.font.size = Pt(12)
    placeholder = doc.add_paragraph("[[TOC]]")
    placeholder.paragraph_format.first_line_indent = Cm(0)
    doc.add_page_break()


def add_section(doc: Document, heading: str, body: str, level: int = 1) -> None:
    doc.add_heading(heading, level=level)
    add_markdown_block(doc, body)


def main() -> None:
    intro = sections(INTRO)
    theory = sections(THEORY)
    development = sections(DEVELOPMENT)
    summary = sections(SUMMARY)

    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_front_matter(doc, summary)

    doc.add_heading("1 INTRODUÇÃO", level=1)
    for number, name in enumerate(("Contextualização e justificativa", "Problema de pesquisa", "Objetivo geral", "Objetivos específicos", "Delimitação"), start=1):
        add_section(doc, f"1.{number} {name}", intro[name], level=2)

    doc.add_heading("2 FUNDAMENTAÇÃO TEÓRICA", level=1)
    theory_names = [
        "Sistemas de informação aplicados à gestão comercial",
        "Lead scoring e priorização de oportunidades",
        "Aprendizado de máquina no pipeline de vendas",
        "Explicabilidade e participação humana",
        "Assistente em linguagem natural com consultas controladas",
        "Síntese e relação com a solução proposta",
    ]
    for number, name in enumerate(theory_names, start=1):
        add_section(doc, f"2.{number} {name}", theory[name], level=2)

    add_section(doc, "3 METODOLOGIA", development["Método"], level=1)

    doc.add_heading("4 DESENVOLVIMENTO E AVALIAÇÃO PARCIAL", level=1)
    development_names = [
        "Base de dados",
        "Arquitetura inicial",
        "Experimento preditivo inicial",
        "Ameaças à validade",
        "Estado do desenvolvimento",
    ]
    for number, name in enumerate(development_names, start=1):
        add_section(doc, f"4.{number} {name}", development[name], level=2)
        if name == "Arquitetura inicial":
            add_figure(
                doc,
                ARCHITECTURE_FIGURE,
                "Figura 1 - Arquitetura da versão parcial e componentes planejados",
                "Diagrama da arquitetura com usuário, dashboard, API, dados, priorização auditável e evolução planejada.",
            )
        elif name == "Estado do desenvolvimento":
            add_figure(
                doc,
                DASHBOARD_FIGURE,
                "Figura 2 - Dashboard conectado à base de demonstração",
                "Captura do dashboard com indicadores, funil comercial e navegação do protótipo.",
            )

    doc.add_heading("5 CONSIDERAÇÕES PARCIAIS", level=1)
    add_markdown_block(
        doc,
        "A versão atual demonstra a importação reproduzível dos dados, o cálculo de indicadores, uma API local e um dashboard responsivo. A avaliação temporal não apresentou evidência suficiente para publicar as saídas dos modelos como probabilidades comerciais confiáveis. Por isso, a fila utiliza critérios auditáveis enquanto a modelagem permanece como experimento acadêmico.\n\n"
        "Os resultados ainda não permitem afirmar impacto sobre conversão ou receita. As próximas etapas incluem autenticação, envio controlado de arquivos, migração para PostgreSQL, avaliação com usuários e integração do assistente de linguagem natural por funções testadas. A conclusão final dependerá dessas etapas e da revisão dos requisitos institucionais.",
    )

    doc.add_heading("REFERÊNCIAS", level=1)
    references = []
    for source in (theory["Referências"], development["Referências técnicas consultadas nesta etapa"]):
        references.extend([clean_inline(item) for item in re.split(r"\n\s*\n", source) if item.strip()])
    seen: set[str] = set()
    for reference in references:
        key = reference.casefold()
        if key in seen:
            continue
        seen.add(key)
        doc.add_paragraph(reference, style="Reference")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
