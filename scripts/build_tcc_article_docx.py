from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "tcc_texto" / "TCC_Leonardo_Caseri_Parcial_ABNT.docx"
INTRO = ROOT / "tcc_texto" / "INTRODUCAO_PARCIAL.md"
THEORY = ROOT / "tcc_texto" / "FUNDAMENTACAO_TEORICA_PARCIAL.md"
DEVELOPMENT = ROOT / "tcc_texto" / "DESENVOLVIMENTO_PARCIAL.md"
SUMMARY = ROOT / "tcc_texto" / "RESUMO_PARCIAL.md"
DASHBOARD_FIGURE = ROOT / "artifacts" / "figures" / "dashboard-overview.png"


def sections(path: Path) -> dict[str, str]:
    text = path.read_text(encoding="utf-8")
    result: dict[str, str] = {}
    current: str | None = None
    lines: list[str] = []
    for line in text.splitlines():
        if line.startswith("## "):
            if current is not None:
                result[current] = "\n".join(lines).strip()
            current = line[3:].strip()
            lines = []
        elif current is not None:
            lines.append(line)
    if current is not None:
        result[current] = "\n".join(lines).strip()
    return result


def set_font(run, name: str = "Times New Roman", size: float | None = None) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    for theme_attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{theme_attribute}"), None)
    if size is not None:
        run.font.size = Pt(size)


def clean_text(text: str) -> str:
    return text.replace("—", "-").replace("–", "-").strip()


def add_inline_runs(paragraph, text: str, size: float | None = None) -> None:
    text = clean_text(text)
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_font(run, size=size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(token[1:-1])
        set_font(run, size=size)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_font(run, size=size)


def set_style_font(style, size: float, bold: bool = False) -> None:
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attribute}"), "Times New Roman")
    for theme_attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{theme_attribute}"), None)


def remove_style_border(style) -> None:
    paragraph_properties = style._element.get_or_add_pPr()
    border = paragraph_properties.find(qn("w:pBdr"))
    if border is not None:
        paragraph_properties.remove(border)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D9D9D9", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, value: int = 85) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge in ("top", "left", "bottom", "right"):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def mark_row_as_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        node = OxmlElement("w:tblHeader")
        node.set(qn("w:val"), "true")
        tr_pr.append(node)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    set_font(run, size=10)
    for tag, field_type, content in (
        ("w:fldChar", "begin", None),
        ("w:instrText", None, " PAGE "),
        ("w:fldChar", "separate", None),
        ("w:t", None, "1"),
        ("w:fldChar", "end", None),
    ):
        node = OxmlElement(tag)
        if field_type:
            node.set(qn("w:fldCharType"), field_type)
        if content is not None:
            node.text = content
        run._r.append(node)


def set_picture_alt_text(paragraph, description: str) -> None:
    for doc_pr in paragraph._p.xpath(".//wp:docPr"):
        doc_pr.set("descr", description)


def add_caption(doc: Document, text: str, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_font(run, size=10)


def add_source(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_font(run, size=10)


def add_figure(doc: Document, path: Path, caption: str, description: str, width: float = 15.2) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    add_caption(doc, caption)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_picture(str(path), width=Cm(width))
    set_picture_alt_text(paragraph, description)
    add_source(doc, "Fonte: elaboração própria (2026).")


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if rows and rows[0] and rows[0][0] == "Modelo":
        add_caption(doc, "Tabela 1 - Desempenho dos procedimentos no conjunto de teste")
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    mark_row_as_header(table.rows[0])
    widths = [Cm(7.0), Cm(4.0), Cm(4.0)] if len(rows[0]) == 3 else [Cm(15.0 / len(rows[0]))] * len(rows[0])
    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            cell.width = widths[col_index]
            cell.text = clean_text(re.sub(r"[*`]", "", value))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, "404040")
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F2F2F2")
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_font(run, size=9)
                    if row_index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    add_source(doc, "Fonte: elaboração própria, a partir dos resultados do experimento (2026).")


def add_markdown_block(doc: Document, block: str) -> None:
    lines = block.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line or line.startswith(">"):
            index += 1
            continue
        if line.startswith("| ") and index + 1 < len(lines) and "---" in lines[index + 1]:
            table_lines = [line]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = [[value.strip() for value in row.strip("|").split("|")] for row in table_lines]
            add_table(doc, rows)
            continue
        if line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.first_line_indent = Cm(0)
            add_inline_runs(paragraph, line[2:])
            index += 1
            continue
        paragraph_lines = [line]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or candidate.startswith("- ") or candidate.startswith("| "):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline_runs(paragraph, " ".join(paragraph_lines))


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    normal = doc.styles["Normal"]
    set_style_font(normal, 12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)

    for style_name, size in (("Title", 14), ("Heading 1", 12), ("Heading 2", 12)):
        style = doc.styles[style_name]
        set_style_font(style, size, bold=True)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    remove_style_border(doc.styles["Title"])
    doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.styles["Heading 2"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    list_style = doc.styles["List Bullet"]
    set_style_font(list_style, 12)
    list_style.paragraph_format.left_indent = Cm(1.25)
    list_style.paragraph_format.first_line_indent = Cm(-0.5)
    list_style.paragraph_format.line_spacing = 1.5
    list_style.paragraph_format.space_after = Pt(0)

    for style_name, left_indent in (("TOC 1", 0), ("TOC 2", 0.75)):
        if style_name not in [style.name for style in doc.styles]:
            toc_style = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            toc_style = doc.styles[style_name]
        set_style_font(toc_style, 12)
        toc_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        toc_style.paragraph_format.left_indent = Cm(left_indent)
        toc_style.paragraph_format.first_line_indent = Cm(0)
        toc_style.paragraph_format.line_spacing = 1.0
        toc_style.paragraph_format.space_after = Pt(0)
        toc_style.paragraph_format.tab_stops.add_tab_stop(
            Cm(15.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )

    if "Reference" not in [style.name for style in doc.styles]:
        reference = doc.styles.add_style("Reference", WD_STYLE_TYPE.PARAGRAPH)
    else:
        reference = doc.styles["Reference"]
    set_style_font(reference, 12)
    reference.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    reference.paragraph_format.first_line_indent = Cm(0)
    reference.paragraph_format.line_spacing = 1.0
    reference.paragraph_format.space_after = Pt(12)

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def add_front_matter(doc: Document, summary: dict[str, str]) -> None:
    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(9)
    title.add_run("DESENVOLVIMENTO DE UM ASSISTENTE INTELIGENTE BASEADO EM IA PARA ANÁLISE DE DADOS EMPRESARIAIS")

    english_title = doc.add_paragraph()
    english_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    english_title.paragraph_format.first_line_indent = Cm(0)
    english_title.paragraph_format.line_spacing = 1.0
    english_title.paragraph_format.space_after = Pt(12)
    run = english_title.add_run("DEVELOPMENT OF AN AI-BASED INTELLIGENT ASSISTANT FOR BUSINESS DATA ANALYSIS")
    run.bold = True
    set_font(run, size=12)

    for text, bold in (("Leonardo Caseri", True), ("Felipe Diniz Dalilo", False)):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        run.bold = bold
        set_font(run, size=11)

    affiliation = doc.add_paragraph()
    affiliation.alignment = WD_ALIGN_PARAGRAPH.LEFT
    affiliation.paragraph_format.first_line_indent = Cm(0)
    affiliation.paragraph_format.line_spacing = 1.0
    affiliation.paragraph_format.space_before = Pt(6)
    affiliation.paragraph_format.space_after = Pt(9)
    run = affiliation.add_run(
        "Graduando do curso de Sistemas de Informação da Universidade de Araraquara - UNIARA. "
        "Araraquara, SP. E-mail: lcaseri@uniara.edu.br.\n"
        "Orientador e docente do curso de Sistemas de Informação da Universidade de Araraquara - UNIARA. "
        "Araraquara, SP. E-mail: fddallilo@uniara.edu.br."
    )
    set_font(run, size=9)

    for title_text, key, keywords_label in (
        ("RESUMO", "Resumo", "Palavras-chave"),
        ("ABSTRACT", "Abstract", "Keywords"),
    ):
        heading = doc.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
        heading.paragraph_format.first_line_indent = Cm(0)
        heading.paragraph_format.line_spacing = 1.0
        heading.paragraph_format.space_before = Pt(6)
        heading.paragraph_format.space_after = Pt(3)
        run = heading.add_run(title_text)
        run.bold = True
        set_font(run, size=12)

        body, keywords = summary[key].rsplit(f"\n\n{keywords_label}:", 1)
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_after = Pt(3)
        add_inline_runs(paragraph, body, size=10)

        keyword_paragraph = doc.add_paragraph()
        keyword_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        keyword_paragraph.paragraph_format.first_line_indent = Cm(0)
        keyword_paragraph.paragraph_format.line_spacing = 1.0
        keyword_paragraph.paragraph_format.space_after = Pt(3)
        label = keyword_paragraph.add_run(f"{keywords_label}: ")
        label.bold = True
        set_font(label, size=10)
        add_inline_runs(keyword_paragraph, keywords.strip(), size=10)


def add_toc(doc: Document) -> None:
    doc.add_page_break()

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.first_line_indent = Cm(0)
    heading.paragraph_format.line_spacing = 1.0
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(12)
    run = heading.add_run("SUMÁRIO")
    run.bold = True
    set_font(run, size=12)

    toc = doc.add_paragraph()
    toc.paragraph_format.first_line_indent = Cm(0)
    toc.paragraph_format.line_spacing = 1.0
    toc.paragraph_format.space_after = Pt(0)
    field_run = toc.add_run()

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-2" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Atualize o sumário no Microsoft Word."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, placeholder, end):
        field_run._r.append(element)
    set_font(field_run, size=12)


def start_textual_section(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    page_numbering = section._sectPr.find(qn("w:pgNumType"))
    if page_numbering is None:
        page_numbering = OxmlElement("w:pgNumType")
        section._sectPr.append(page_numbering)
    page_numbering.set(qn("w:start"), "1")
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.header.paragraphs[0].clear()
    add_page_number(section.header.paragraphs[0])


def add_primary_heading(doc: Document, text: str, first: bool = False, centered: bool = False) -> None:
    if not first and centered:
        doc.add_page_break()
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["Heading 1"]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.add_run(text)


def add_section(doc: Document, heading: str, body: str, level: int = 2) -> None:
    doc.add_heading(heading, level=level)
    add_markdown_block(doc, body)


def add_references(doc: Document, theory: dict[str, str], development: dict[str, str]) -> None:
    references: list[str] = []
    for source in (theory["Referências"], development["Referências técnicas consultadas nesta etapa"]):
        references.extend(item.strip() for item in re.split(r"\n\s*\n", source) if item.strip())
    unique: dict[str, str] = {}
    for reference in references:
        key = clean_text(re.sub(r"[*`]", "", reference)).casefold()
        unique.setdefault(key, reference)
    for reference in sorted(unique.values(), key=lambda value: re.sub(r"[*`]", "", value).casefold()):
        paragraph = doc.add_paragraph(style="Reference")
        add_inline_runs(paragraph, reference)


def main() -> None:
    intro = sections(INTRO)
    theory = sections(THEORY)
    development = sections(DEVELOPMENT)
    summary = sections(SUMMARY)

    doc = Document()
    doc.core_properties.title = "Desenvolvimento de um assistente inteligente baseado em IA para análise de dados empresariais"
    doc.core_properties.author = "Leonardo Caseri"
    doc.core_properties.subject = "Trabalho de Conclusão de Curso em formato de artigo científico"
    configure_document(doc)
    add_front_matter(doc, summary)
    add_toc(doc)
    start_textual_section(doc)

    add_primary_heading(doc, "1 INTRODUÇÃO", first=True)
    for number, name in enumerate(
        ("Contextualização e justificativa", "Problema de pesquisa", "Objetivo geral", "Objetivos específicos", "Delimitação"),
        start=1,
    ):
        add_section(doc, f"1.{number} {name}", intro[name])

    add_primary_heading(doc, "2 FUNDAMENTAÇÃO TEÓRICA")
    for number, name in enumerate(
        (
            "Sistemas de informação aplicados à gestão comercial",
            "Lead scoring e priorização de oportunidades",
            "Aprendizado de máquina no pipeline de vendas",
            "Explicabilidade e participação humana",
            "Assistente em linguagem natural com consultas controladas",
            "Síntese e relação com a solução proposta",
        ),
        start=1,
    ):
        add_section(doc, f"2.{number} {name}", theory[name])

    add_primary_heading(doc, "3 METODOLOGIA")
    method_names = (
        "Caracterização da pesquisa",
        "Coleta e preparação dos dados",
        "Desenvolvimento e verificação funcional",
        "Avaliação preditiva",
        "Aspectos éticos e limites do método",
    )
    for number, name in enumerate(method_names, start=1):
        add_section(doc, f"3.{number} {name}", development[name])

    add_primary_heading(doc, "4 DESENVOLVIMENTO DO SISTEMA E RESULTADOS")
    development_names = (
        "Base de dados",
        "Arquitetura e organização do código",
        "Importação e persistência dos dados",
        "API com autenticação e isolamento dos dados",
        "Interface de indicadores e fila de revisão",
        "Experimento preditivo inicial",
        "Assistente analítico no escopo final",
        "Verificação funcional do protótipo",
        "Ameaças à validade",
        "Relação entre a implementação e o protótipo final",
    )
    for number, name in enumerate(development_names, start=1):
        add_section(doc, f"4.{number} {name}", development[name])
        if name == "Interface de indicadores e fila de revisão":
            add_figure(
                doc,
                DASHBOARD_FIGURE,
                "Figura 1 - Dashboard conectado à base de demonstração",
                "Captura do dashboard com indicadores, funil comercial e fila de oportunidades.",
            )

    add_primary_heading(doc, "5 DISCUSSÃO DOS RESULTADOS")
    discussion_names = (
        "Contribuições técnicas do protótipo avaliado",
        "Interpretação do experimento preditivo",
        "Implicações para uso comercial e continuidade",
    )
    for number, name in enumerate(discussion_names, start=1):
        add_section(doc, f"5.{number} {name}", development[name])

    add_primary_heading(doc, "REFERÊNCIAS", centered=True)
    add_references(doc, theory, development)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
